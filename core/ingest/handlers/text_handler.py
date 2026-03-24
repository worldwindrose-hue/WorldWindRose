"""
ROSA OS — Text Handler.

Handles TXT, MD, JSON, EPUB, DOCX, CSV, XLSX, and plain code files.
"""

from __future__ import annotations

import csv
import json
import logging
from pathlib import Path

from core.ingest.handlers.base import BaseHandler, IngestResult

logger = logging.getLogger("rosa.ingest.text")


class TextHandler(BaseHandler):
    """Ingest text-based files into the knowledge graph."""

    async def process(self, job) -> IngestResult:
        source = job.source
        self.update_progress(job, 5, "Читаю файл...")
        try:
            path = Path(source)
            ext = path.suffix.lower()
            text = self._read(path, ext)

            if not text.strip():
                raise ValueError("Файл пустой или не содержит текста")

            self.update_progress(job, 60, "Сохраняю в граф знаний...")
            chunks = self.chunk(text)
            tags = [ext.lstrip(".") or "text"]
            nodes = await self.save_to_graph(chunks, source=source, tags=tags,
                                              extra_meta={"filename": path.name, "ext": ext})

            self.update_progress(job, 100)
            return IngestResult(
                type="text",
                source=source,
                nodes_created=nodes,
                chunks=len(chunks),
                summary=f"✅ {path.name}: {len(chunks)} чанков → {nodes} узлов",
                metadata={"filename": path.name, "ext": ext},
            )
        except Exception as exc:
            logger.error("Text ingest failed: %s", exc)
            raise

    def _read(self, path: Path, ext: str) -> str:
        if ext in (".txt", ".md", ".rst", ".log", ".py", ".js", ".ts",
                   ".go", ".rs", ".java", ".cpp", ".c", ".cs", ".rb",
                   ".php", ".swift", ".kt", ".sql", ".sh", ".toml",
                   ".yaml", ".yml", ".ini", ".cfg", ".env", ".html", ".xml"):
            return self._read_plain(path)
        elif ext == ".json":
            return self._read_json(path)
        elif ext in (".csv",):
            return self._read_csv(path)
        elif ext in (".xlsx", ".xls"):
            return self._read_excel(path)
        elif ext == ".docx":
            return self._read_docx(path)
        elif ext == ".epub":
            return self._read_epub(path)
        else:
            return self._read_plain(path)

    def _read_plain(self, path: Path) -> str:
        for enc in ("utf-8", "utf-8-sig", "cp1251", "latin-1"):
            try:
                return path.read_text(encoding=enc)
            except (UnicodeDecodeError, LookupError):
                continue
        return path.read_bytes().decode("utf-8", errors="replace")

    def _read_json(self, path: Path) -> str:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            return json.dumps(data, ensure_ascii=False, indent=2)
        except Exception:
            return self._read_plain(path)

    def _read_csv(self, path: Path) -> str:
        lines = []
        try:
            with open(path, newline="", encoding="utf-8-sig") as f:
                reader = csv.reader(f)
                for i, row in enumerate(reader):
                    if i > 5000:
                        lines.append(f"... (файл обрезан после 5000 строк)")
                        break
                    lines.append(" | ".join(row))
        except Exception as exc:
            logger.warning("CSV read failed: %s", exc)
            return self._read_plain(path)
        return "\n".join(lines)

    def _read_excel(self, path: Path) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheets = []
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True, max_row=5000):
                    row_str = " | ".join("" if v is None else str(v) for v in row)
                    if row_str.strip():
                        rows.append(row_str)
                if rows:
                    sheets.append(f"=== {ws.title} ===\n" + "\n".join(rows))
            return "\n\n".join(sheets)
        except ImportError:
            logger.warning("openpyxl not installed")
            return ""
        except Exception as exc:
            logger.warning("Excel read failed: %s", exc)
            return ""

    def _read_docx(self, path: Path) -> str:
        try:
            import docx
            doc = docx.Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            logger.warning("python-docx not installed")
            return ""
        except Exception as exc:
            logger.warning("DOCX read failed: %s", exc)
            return ""

    def _read_epub(self, path: Path) -> str:
        try:
            import ebooklib
            from ebooklib import epub
            from bs4 import BeautifulSoup
            book = epub.read_epub(str(path))
            parts = []
            for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
                soup = BeautifulSoup(item.get_content(), "html.parser")
                text = soup.get_text(separator="\n", strip=True)
                if text:
                    parts.append(text)
            return "\n\n".join(parts)
        except ImportError:
            logger.warning("ebooklib/bs4 not installed")
            return ""
        except Exception as exc:
            logger.warning("EPUB read failed: %s", exc)
            return ""
